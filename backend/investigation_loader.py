"""
Investigation Dashboard Data Loader & Normalizer
Reads CRI and CRIN investigation sheets from the 2026 Investigation workbook.
Provides in-memory caching and clean data models for dashboard and analytics.
"""

import os
import re
import threading
from datetime import datetime, date
from typing import Dict, List, Any, Optional
from openpyxl import load_workbook

# Default workbook path inside repository
DEFAULT_EXCEL_FILENAME = "2026(3) - Investigation Dashboard updated.xlsx"

# Environment variable name for override
ENV_EXCEL_VAR = "INVESTIGATION_EXCEL_FILE"

# Thread safety lock for cache
_CACHE_LOCK = threading.Lock()
_DATA_CACHE: Optional[Dict[str, Any]] = None
_CACHE_TIMESTAMP: Optional[datetime] = None


def get_excel_file_path() -> str:
    """Return the active Excel file path from env var or repository default."""
    env_path = os.environ.get(ENV_EXCEL_VAR)
    if env_path and env_path.strip():
        return env_path.strip()
    
    # Try current directory or parent directory
    base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    local_path = os.path.join(base_dir, DEFAULT_EXCEL_FILENAME)
    if os.path.exists(local_path):
        return local_path
    
    # Fallback to current working directory
    cwd_path = os.path.abspath(DEFAULT_EXCEL_FILENAME)
    if os.path.exists(cwd_path):
        return cwd_path
        
    return local_path



def _clean_str(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, (datetime, date)):
        return val.strftime("%d-%m-%Y")
    s = str(val).strip()
    return "" if s.lower() == "nan" else s


def _as_int(val: Any, default: int = 1) -> int:
    if val in (None, "", "nan", "-"):
        return default
    try:
        return int(float(val))
    except (ValueError, TypeError):
        return default


def _as_float(val: Any, default: float = 0.0) -> float:
    if val in (None, "", "nan", "-"):
        return default
    try:
        return float(val)
    except (ValueError, TypeError):
        return default


MONTH_LOOKUP = {
    "jan": (2026, 1, "2026-01", "Jan 2026"),
    "feb": (2026, 2, "2026-02", "Feb 2026"),
    "mar": (2026, 3, "2026-03", "Mar 2026"),
    "apr": (2026, 4, "2026-04", "Apr 2026"),
    "may": (2026, 5, "2026-05", "May 2026"),
    "jun": (2026, 6, "2026-06", "Jun 2026"),
    "jul": (2026, 7, "2026-07", "Jul 2026"),
    "aug": (2026, 8, "2026-08", "Aug 2026"),
    "sep": (2026, 9, "2026-09", "Sep 2026"),
    "oct": (2026, 10, "2026-10", "Oct 2026"),
    "nov": (2026, 11, "2026-11", "Nov 2026"),
    "dec": (2026, 12, "2026-12", "Dec 2026"),
}


def normalize_month(val: Any) -> Dict[str, str]:
    """Normalize raw month input (strings or datetime) into code and label."""
    if val is None:
        return {"code": "2026-00", "label": "Unknown", "raw": ""}
        
    if isinstance(val, (datetime, date)):
        return {
            "code": val.strftime("%Y-%m"),
            "label": val.strftime("%b %Y"),
            "raw": str(val),
        }
        
    raw_s = str(val).strip()
    s = raw_s.lower()
    for prefix, (yr, m_num, code, label) in MONTH_LOOKUP.items():
        if prefix in s:
            return {"code": code, "label": label, "raw": raw_s}
            
    return {"code": "2026-00", "label": raw_s or "Unknown", "raw": raw_s}


def normalize_customer(val: Any) -> str:
    """Standardize customer acronyms without altering meaning."""
    clean = _clean_str(val).upper().replace(" ", "")
    if clean in ("TCL", "TATACOMMERCIAL"):
        return "TCL"
    if clean in ("M&M", "MM", "MAHINDRA"):
        return "M&M"
    if clean in ("TML", "TATAMOTORS"):
        return "TML"
    if clean in ("A.L", "AL", "ASHOKLEYLAND"):
        return "AL"
    if clean in ("VECV", "VOLVOEICHER"):
        return "VECV"
    if clean in ("ISUZU",):
        return "ISUZU"
    if clean in ("SMLI", "SMLISUZU"):
        return "SMLI"
    if clean in ("TTF",):
        return "TTF"
    if clean in ("KOEL", "KIRLOSKAR"):
        return "KOEL"
    if clean in ("CNHI",):
        return "CNHI"
    if clean in ("FUSO",):
        return "FUSO"
    return _clean_str(val)


def normalize_complaint_type(val: Any) -> str:
    """Standardize complaint type (0 Km vs Field vs AIS)."""
    s = _clean_str(val).strip().upper()
    if "0" in s or "KM" in s or "0KM" in s:
        return "0 Km"
    if "FIELD" in s:
        return "Field"
    if "AIS" in s:
        return "AIS"
    return _clean_str(val) or "Unspecified"


def classify_defect_finding(finding_str: str) -> Dict[str, str]:
    """
    Derives Investigation Status, Outcome, and Defect Category.
    Rules:
      - Blank finding -> Status: Pending, Outcome: Pending, Category: Pending / In Analysis
      - Contains 'conforming' -> Status: Completed, Outcome: Conforming to Specifications, Category: Conforming to Specifications
      - Contains defect keywords -> Status: Completed, Outcome: Defect Identified, Category: <Specific Failure Mode>
    """
    s = finding_str.strip()
    if not s:
        return {
            "status": "Pending",
            "outcome": "Pending",
            "category": "Pending / In Analysis",
        }
        
    s_lower = s.lower()
    if any(k in s_lower for k in ("conforming", "confroming", "conforimng")):
        return {
            "status": "Completed",
            "outcome": "Conforming to Specifications",
            "category": "Conforming to Specifications",
        }
        
    outcome = "Defect Identified"
    
    if "z hole" in s_lower or "z-hole" in s_lower or "valve piece z" in s_lower:
        category = "Particle at Z-Hole"
    elif "nozzle seat" in s_lower or "nozzle tip" in s_lower or "nozzle needle" in s_lower or "nozzle body" in s_lower:
        category = "Nozzle Area Particle / Defect"
    elif "contamination" in s_lower or "external particle" in s_lower or "liquid" in s_lower:
        category = "Fuel Contamination"
    elif "magnet" in s_lower or "coil" in s_lower or "short to ground" in s_lower:
        category = "Magnet / Coil Defect"
    elif "crack" in s_lower or "stress corrosion" in s_lower:
        category = "Crack / Stress Corrosion"
    elif "tempering" in s_lower or "tampering" in s_lower or "damage" in s_lower:
        category = "External Tampering / Handling"
    else:
        category = "Other Mechanical / Material Finding"
        
    return {
        "status": "Completed",
        "outcome": outcome,
        "category": category,
    }


def parse_sheet_rows(sheet, product_class: str) -> List[Dict[str, Any]]:
    """Parse raw openpyxl worksheet rows for CRI or CRIN into structured dictionaries."""
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
        
    # First row is header
    header_row = rows[0]
    col_map = {}
    for idx, cell in enumerate(header_row):
        if cell is not None:
            clean_h = str(cell).strip()
            col_map[clean_h] = idx
            
    records = []
    
    for row_idx, row in enumerate(rows[1:], start=2):
        # Ignore completely empty rows
        if not any(c is not None and str(c).strip() != "" for c in row):
            continue
            
        def get_val(*col_names):
            for name in col_names:
                if name in col_map:
                    idx = col_map[name]
                    if idx < len(row):
                        return row[idx]
            return None

        sr_no = get_val("Sr.No", "Sr. No", "Sr No")
        complaint_type_raw = get_val("Complaint type", "Complaint Type")
        mileage_raw = get_val("Mileage", "Milage (KM)", "Milage")
        qc_number = _clean_str(get_val("QC Number", "QC No"))
        month_raw = get_val("Month")
        injector_number = _clean_str(get_val("Injector Number", " Injector Number"))
        customer_raw = get_val("Customer")
        complaint_symptom = _clean_str(get_val("Complaint"))
        complaint_no = _clean_str(get_val("Customer Complaint Number", "Complaint Number"))
        esn = _clean_str(get_val("ESN"))
        qty_raw = get_val("Qty")
        serial_no = _clean_str(get_val("Serial No.", "Serial No"))
        mfd = _clean_str(get_val("MFD"))
        for_analysis = _clean_str(get_val("For analysis", "For Analysis"))
        investigation_finding = _clean_str(get_val("Investigation Finding"))
        finding_raw = _clean_str(get_val("Finding"))
        particle_location = _clean_str(get_val("Particle Location"))
        responsibility = _clean_str(get_val("Responsibility"))
        part_received_date = _clean_str(get_val("Part Received Date", "Parts Received Date"))
        report_shared_date = _clean_str(get_val("Investigation report shared date"))
        investigation_days = _clean_str(get_val("Investigation days"))
        pending_reason = _clean_str(get_val("Investigation Pending Reason"))
        complaint_mode = _clean_str(get_val("Complaint Mode"))
        mfg_plant = _clean_str(get_val("Mfg Plant"))
        product_class_val = _clean_str(get_val("Product Class")) or product_class

        month_info = normalize_month(month_raw)
        customer_clean = normalize_customer(customer_raw)
        complaint_type_clean = normalize_complaint_type(complaint_type_raw)
        qty_val = _as_int(qty_raw, default=1)
        
        # Classification
        classification = classify_defect_finding(investigation_finding)

        record = {
            "id": f"{product_class}_{row_idx}",
            "row_number": row_idx,
            "product_class": product_class,
            "sr_no": sr_no,
            "customer_complaint_no": complaint_no,
            "complaint_type": complaint_type_clean,
            "customer": customer_clean,
            "month_raw": _clean_str(month_raw),
            "month_code": month_info["code"],
            "month_label": month_info["label"],
            "injector_number": injector_number,
            "complaint": complaint_symptom,
            "qty": qty_val,
            "serial_no": serial_no,
            "mfd": mfd,
            "mileage": _clean_str(mileage_raw),
            "qc_number": qc_number if qc_number != "-" else "",
            "esn": esn if esn != "-" else "",
            "investigation_finding": investigation_finding,
            "finding": finding_raw,
            "particle_location": particle_location,
            "responsibility": responsibility,
            "mfg_plant": mfg_plant,
            "for_analysis": for_analysis,
            "part_received_date": part_received_date,
            "report_shared_date": report_shared_date,
            "investigation_days": investigation_days,
            "pending_reason": pending_reason,
            "complaint_mode": complaint_mode,
            "status": classification["status"],
            "outcome": classification["outcome"],
            "defect_category": classification["category"],
        }
        records.append(record)
        
    return records


def load_investigation_workbook(excel_path: Optional[str] = None) -> Dict[str, Any]:
    """Load and normalize workbook records from CRI and CRIN worksheets."""
    path = excel_path or get_excel_file_path()
    
    result = {
        "source_file": path,
        "loaded_at": datetime.now(),
        "error": "",
        "records": [],
        "cri_records": [],
        "crin_records": [],
        "field_claims": [],
        "total_records": 0,
        "months": [],
        "customers": [],
        "complaint_types": [],
        "defect_categories": [],
        "injector_numbers": [],
    }
    
    if not os.path.exists(path):
        result["error"] = f"Investigation workbook not found at: {path}"
        return result
        
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
        
        cri_records = []
        crin_records = []
        
        if "CRI" in wb.sheetnames:
            cri_records = parse_sheet_rows(wb["CRI"], "CRI")
        else:
            result["error"] = "Sheet 'CRI' not found in workbook."
            
        if "CRIN" in wb.sheetnames:
            crin_records = parse_sheet_rows(wb["CRIN"], "CRIN")
        else:
            if not result["error"]:
                result["error"] = "Sheet 'CRIN' not found in workbook."

        wb.close()
        
        all_records = cri_records + crin_records
        
        # Populate unique filter option lists
        months_set = {}
        for r in all_records:
            if r["month_code"] != "2026-00" and r["month_label"]:
                months_set[r["month_code"]] = r["month_label"]
        sorted_months = [
            {"code": code, "label": label}
            for code, label in sorted(months_set.items())
        ]
        
        customers = sorted(list({r["customer"] for r in all_records if r["customer"]}))
        complaint_types = sorted(list({r["complaint_type"] for r in all_records if r["complaint_type"]}))
        defect_categories = sorted(list({r["defect_category"] for r in all_records if r["defect_category"]}))
        injector_numbers = sorted(list({r["injector_number"] for r in all_records if r["injector_number"]}))

        result.update({
            "records": all_records,
            "cri_records": cri_records,
            "crin_records": crin_records,
            "total_records": len(all_records),
            "months": sorted_months,
            "customers": customers,
            "complaint_types": complaint_types,
            "defect_categories": defect_categories,
            "injector_numbers": injector_numbers,
        })
        return result
        
    except Exception as e:
        result["error"] = f"Error reading investigation Excel file: {str(e)}"
        return result


def get_cached_investigation_data(force_reload: bool = False) -> Dict[str, Any]:
    """Retrieve investigation data from thread-safe memory cache or load if needed."""
    global _DATA_CACHE, _CACHE_TIMESTAMP
    
    with _CACHE_LOCK:
        if _DATA_CACHE is None or force_reload:
            _DATA_CACHE = load_investigation_workbook()
            _CACHE_TIMESTAMP = datetime.now()
        return _DATA_CACHE


def reload_investigation_data() -> Dict[str, Any]:
    """Explicitly invalidate cache and reload fresh from the workbook."""
    return get_cached_investigation_data(force_reload=True)
