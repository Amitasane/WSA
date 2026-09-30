import os
import json
import hashlib
from datetime import datetime, date
from typing import Dict, List, Any, Optional
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from backend.alerts.models import DataRecordSnapshot

PARTICLE_SHEET = "Particle Summary"
PARTICLE_HEADERS = [
    "Date",
    "Shift",
    "Auditor",
    "Type",
    "Customer",
    "Station",
    "Defective parts checked",
    "Particle inside nozzle",
    "Particle in Z hole",
    "Particle in A hole",
    "Total particle found",
    "Rejection Date",
    "Rejection Shift",
    "Metallic/Non metallic",
    "DA/Fresh",
    "Chemistry",
    "Size",
]

BOSCH_DFS_PARTICLE_PATH = (
    r"\\bosch.com\dfsrb\DfsIN\LOC\Na\DS\QMM\02_Projects\04_QMM3_Common"
    r"\08_Projects_GAs\03_LAC\2024 - Snehal Yelmalle\CRIN Internal Rejection Analysis"
    r"\CRIN Line rejection analysis updated.xlsx"
)


def _clean_val(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, (datetime, date)):
        return val.strftime("%d-%m-%Y")
    s = str(val).strip()
    return "" if s.lower() == "nan" else s


def _iso_date(val: Any) -> str:
    if val in (None, ""):
        return ""
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    text = str(val).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return ""


def _as_num(val: Any) -> float:
    if val in (None, "", "-", "~"):
        return 0.0
    try:
        return float(val)
    except (ValueError, TypeError):
        try:
            return float(str(val).replace(",", "").strip())
        except (ValueError, TypeError):
            return 0.0


def find_particle_workbook_path() -> Optional[str]:
    """Locate the WSA 'CRIN Line rejection analysis updated.xlsx' workbook."""
    env_path = os.environ.get("WSA_EXCEL_FILE", "").strip()
    if env_path and os.path.exists(env_path):
        return env_path

    if os.path.exists(BOSCH_DFS_PARTICLE_PATH):
        return BOSCH_DFS_PARTICLE_PATH

    base_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
    candidates = [
        os.path.join(base_dir, "CRIN Line rejection analysis updated.xlsx"),
        os.path.abspath("CRIN Line rejection analysis updated.xlsx"),
    ]
    for c in candidates:
        if os.path.exists(c):
            return c
    return None


def load_particle_summary_records(excel_path: Optional[str] = None) -> List[Dict[str, Any]]:
    """
    Loads and normalizes records from the WSA 'CRIN Line rejection analysis updated.xlsx'
    ('Particle Summary' sheet), automatically skipping bottom pivot/summary rows that lack a valid Date.
    """
    path = excel_path or find_particle_workbook_path()
    if not path or not os.path.exists(path):
        return []

    try:
        wb = load_workbook(path, data_only=True, read_only=True)
        if PARTICLE_SHEET not in wb.sheetnames:
            wb.close()
            return []

        sheet = wb[PARTICLE_SHEET]
        raw_rows = sheet.iter_rows(values_only=True)
        header_row = next(raw_rows, None)
        if not header_row:
            wb.close()
            return []

        header_map = {_clean_val(val): idx for idx, val in enumerate(header_row)}
        if any(h not in header_map for h in PARTICLE_HEADERS):
            wb.close()
            return []

        records: List[Dict[str, Any]] = []
        for row_idx, raw in enumerate(raw_rows, start=2):
            values = [
                raw[header_map[h]] if header_map[h] < len(raw) else None
                for h in PARTICLE_HEADERS
            ]
            if not any(_clean_val(v) for v in values):
                continue

            # Skip pivot table summary rows at the bottom of the sheet where Date is empty/non-date
            date_iso = _iso_date(values[0])
            if not date_iso:
                continue

            rec: Dict[str, Any] = {
                h: _clean_val(values[i]) for i, h in enumerate(PARTICLE_HEADERS)
            }
            rec["row_number"] = row_idx
            rec["data_source"] = "PARTICLE_SUMMARY"
            rec["_date_iso"] = date_iso
            rec["_rejection_date_iso"] = _iso_date(values[11])
            rec["_parts"] = _as_num(values[6])
            rec["_nozzle_particle"] = _as_num(values[7])
            rec["_zhole_particle"] = _as_num(values[8])
            rec["_ahole_particle"] = _as_num(values[9])
            rec["_total_particle"] = _as_num(values[10])

            particle_val = values[10]
            if particle_val in (None, ""):
                rec["status"] = "PENDING"
            else:
                rec["status"] = "OK" if rec["_total_particle"] == 0 else "NOK"

            locs = []
            if rec["_zhole_particle"] > 0:
                locs.append("Z hole")
            if rec["_nozzle_particle"] > 0:
                locs.append("Inside Nozzle")
            if rec["_ahole_particle"] > 0:
                locs.append("A hole")
            rec["particle_location"] = (
                ", ".join(locs)
                if locs
                else ("Particle observed" if rec["_total_particle"] > 0 else "No Particle found")
            )

            records.append(rec)

        wb.close()
        return records
    except Exception:
        return []


def normalize_unified_record(raw: Dict[str, Any]) -> Dict[str, Any]:
    """
    Normalizes a WSA Particle Summary record into a canonical dictionary
    supporting all rule conditions and Bosch MIS email template variables.
    """
    date_iso = raw.get("_date_iso") or _iso_date(raw.get("Date"))
    parts = _as_num(raw.get("_parts", raw.get("Defective parts checked", 0)))
    total_p = _as_num(raw.get("_total_particle", raw.get("Total particle found", 0)))
    status = (raw.get("status") or ("NOK" if total_p > 0 else "OK")).upper()
    nok_count = 1 if status == "NOK" else 0

    locs = raw.get("particle_location")
    if not locs:
        loc_parts = []
        if _as_num(raw.get("Particle in Z hole")) > 0:
            loc_parts.append("Z hole")
        if _as_num(raw.get("Particle inside nozzle")) > 0:
            loc_parts.append("Inside Nozzle")
        if _as_num(raw.get("Particle in A hole")) > 0:
            loc_parts.append("A hole")
        locs = (
            ", ".join(loc_parts)
            if loc_parts
            else ("Particle observed" if total_p > 0 else "No Particle found")
        )

    return {
        "record_key": f"PS_{date_iso}_{raw.get('Shift','')}_{raw.get('Station','')}_{raw.get('Type','')}_{raw.get('row_number', '')}",
        "data_source": "PARTICLE_SUMMARY",
        "date": raw.get("Date") or date_iso,
        "date_iso": date_iso,
        "shift": raw.get("Shift", ""),
        "station": raw.get("Station", ""),
        "auditor": raw.get("Auditor", ""),
        "customer": raw.get("Customer", ""),
        "injector_type": raw.get("Type", ""),
        "parts_checked": parts,
        "status": status,
        "nok_count": nok_count,
        "total_particles": total_p,
        "particle_nozzle": _as_num(raw.get("Particle inside nozzle", 0)),
        "particle_zhole": _as_num(raw.get("Particle in Z hole", 0)),
        "particle_ahole": _as_num(raw.get("Particle in A hole", 0)),
        "particle_location": locs,
        "rejection_date": raw.get("Rejection Date", ""),
        "rejection_shift": raw.get("Rejection Shift", ""),
        "metallic": raw.get("Metallic/Non metallic", ""),
        "da_fresh": raw.get("DA/Fresh", ""),
        "chemistry": raw.get("Chemistry", ""),
        "size": raw.get("Size", ""),
        "defect_category": "Particle Contamination" if total_p > 0 else "Conforming",
        "investigation_finding": (
            f"Observed {int(total_p)} particle(s) at {locs}"
            if total_p > 0
            else "No abnormality was observed"
        ),
        "investigation_days": 0.0,
        "raw": raw,
    }


def compute_record_hash(norm_rec: Dict[str, Any]) -> str:
    """Compute deterministic SHA-256 hash of a normalized Particle Summary record."""
    payload = {
        "status": norm_rec.get("status"),
        "parts_checked": norm_rec.get("parts_checked"),
        "total_particles": norm_rec.get("total_particles"),
        "particle_location": norm_rec.get("particle_location"),
        "station": norm_rec.get("station"),
        "customer": norm_rec.get("customer"),
        "rejection_date": norm_rec.get("rejection_date"),
        "chemistry": norm_rec.get("chemistry"),
        "size": norm_rec.get("size"),
    }
    encoded = json.dumps(payload, sort_keys=True).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def detect_and_sync_changes(
    db: Session,
    normalized_records: List[Dict[str, Any]],
    seed_as_baseline_if_empty: bool = True,
) -> Dict[str, Any]:
    """
    Compares current CRIN Line rejection analysis Excel records against DataRecordSnapshot table.
    """
    existing_count = db.query(DataRecordSnapshot).count()
    is_initial_seed = seed_as_baseline_if_empty and (existing_count == 0)

    existing_map: Dict[str, DataRecordSnapshot] = {
        snap.record_key: snap for snap in db.query(DataRecordSnapshot).all()
    }

    added: List[Dict[str, Any]] = []
    updated: List[Dict[str, Any]] = []
    status_changed_to_nok: List[Dict[str, Any]] = []
    now = datetime.utcnow()

    for rec in normalized_records:
        r_key = rec["record_key"]
        r_hash = compute_record_hash(rec)
        rec["_row_hash"] = r_hash

        snap = existing_map.get(r_key)
        if snap is None:
            new_snap = DataRecordSnapshot(
                record_key=r_key,
                data_source="PARTICLE_SUMMARY",
                row_hash=r_hash,
                status=rec["status"],
                record_date_iso=rec.get("date_iso", ""),
                station=rec.get("station", ""),
                first_seen_at=now,
                last_seen_at=now,
                summary_json=json.dumps({
                    "date": rec.get("date"),
                    "station": rec.get("station"),
                    "status": rec.get("status"),
                    "total_particles": rec.get("total_particles"),
                }),
            )
            db.add(new_snap)
            existing_map[r_key] = new_snap
            if not is_initial_seed:
                added.append(rec)
                if rec["status"] == "NOK":
                    status_changed_to_nok.append(rec)
        else:
            snap.last_seen_at = now
            if snap.row_hash != r_hash:
                old_status = snap.status
                snap.row_hash = r_hash
                snap.status = rec["status"]
                snap.station = rec.get("station", "")
                if not is_initial_seed:
                    updated.append(rec)
                    if old_status != "NOK" and rec["status"] == "NOK":
                        status_changed_to_nok.append(rec)

    db.commit()
    return {
        "added": added,
        "updated": updated,
        "status_changed_to_nok": status_changed_to_nok,
        "all_records": normalized_records,
        "is_baseline_seed": is_initial_seed,
    }
