import os
from fastapi import UploadFile, File
from fastapi import FastAPI, Request, Depends, Form, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from fastapi.staticfiles import StaticFiles
from sqlalchemy.orm import Session
from sqlalchemy import func, asc, desc
from datetime import datetime, date
import csv
import io
from openpyxl import load_workbook

from backend.database import engine, Base, get_db
from backend import models
from backend.alerts import models as alert_models
from backend.auth import authenticate_user, create_user, get_current_user

# ================= CODE GENERATORS =================

def generate_record_code(date_str, db):
    date_obj = datetime.strptime(date_str, "%Y-%m-%d")
    date_part = date_obj.strftime("%Y%m%d")

    # Count records created on that date
    count = db.query(models.Record).filter(
        models.Record.date == date_obj.date()
    ).count()

    return f"NPR-{date_part}-{str(count + 1).zfill(4)}"


def generate_case_code(case_date, db):
    date_str = case_date.strftime("%Y%m%d")

    # Get highest existing number for this date
    last_case = db.query(models.Case).filter(
        models.Case.case_code.like(f"NPC-{date_str}-%")
    ).order_by(models.Case.case_code.desc()).first()

    if last_case and last_case.case_code:
        last_number = int(last_case.case_code.split("-")[-1])
        new_number = last_number + 1
    else:
        new_number = 1

    return f"NPC-{date_str}-{str(new_number).zfill(4)}"


"""
def generate_case_code(date_obj, db):
    date_part = date_obj.strftime("%Y%m%d")

    # Count cases created on that date
    count = db.query(models.Case).filter(
        models.Case.created_at == date_obj
    ).count()

    return f"NPC-{date_part}-{str(count + 1).zfill(4)}"
"""


app = FastAPI()

# Serve static files
app.mount("/static", StaticFiles(directory="static"), name="static")

Base.metadata.create_all(bind=engine)

templates = Jinja2Templates(directory="templates")


# ================= LOGIN =================

@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    return templates.TemplateResponse("login.html", {
        "request": request,
        "show_header": True,
        "hide_home_icon": True
    })


@app.post("/login")
def login(request: Request,
          username: str = Form(...),
          password: str = Form(...),
          db: Session = Depends(get_db)):

    user = authenticate_user(db, username, password)

    if not user:
        return templates.TemplateResponse("login.html", {
            "request": request,
            "error": "Invalid username or password.",
            "show_header": True,
            "hide_home_icon": True
        })

    response = RedirectResponse("/", status_code=302)
    response.set_cookie("user_id", str(user.id))
    return response


# ================= REGISTER =================

@app.get("/register", response_class=HTMLResponse)
def register_page(request: Request):
    return templates.TemplateResponse("register.html", {
        "request": request,
        "show_header": True,
        "hide_home_icon": True
    })


@app.post("/register")
def register(request: Request,
             first_name: str = Form(...),
             last_name: str = Form(...),
             username: str = Form(...),
             email: str = Form(...),
             password: str = Form(...),
             confirm_password: str = Form(...),
             department: str = Form(...),
             role: str = Form(...),
             db: Session = Depends(get_db)):

    error = create_user(
        db,
        first_name,
        last_name,
        username,
        email,
        password,
        confirm_password,
        department,
        role
    )

    if error:
        return templates.TemplateResponse("register.html", {
            "request": request,
            "error": error,
            "show_header": True,
            "hide_home_icon": True,
            "form_data": {
                "first_name": first_name,
                "last_name": last_name,
                "username": username,
                "email": email,
                "department": department,
                "role": role
            }
        })

    return RedirectResponse("/login", status_code=302)


# ================= LOGOUT =================

@app.get("/logout")
def logout():
    response = RedirectResponse("/login")
    response.delete_cookie("user_id")
    return response


# ================= HOME =================

@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    try:
        user = get_current_user(request)
    except:
        return RedirectResponse("/login")

    return templates.TemplateResponse("home.html", {
        "request": request,
        "user": user,
        "show_header": True
    })


# ================= EXCEL DATA SOURCE =================

EXCEL_SHEET = "Particle Summary"

EXCEL_HEADERS = [
    "Date",
    "Shift",
    "Auditor",
    "Type",
    "Customer",
    "Station",
    "Defective parts checked",
    "Particle inside nozzle",
    "Particle in Z hole",
    "Particle in A hole",
    "Total particle found",
    "Rejection Date",
    "Rejection Shift",
    "Metallic/Non metallic",
    "DA/Fresh",
    "Chemistry",
    "Size",
]

# Central Bosch Excel is the live source for WSA record/quality data.
# WSA_EXCEL_FILE can be set on a workstation for local testing without changing this default.
EXCEL_FILE = os.environ.get(
    "WSA_EXCEL_FILE",
    r"\\bosch.com\dfsrb\DfsIN\LOC\Na\DS\QMM\02_Projects\04_QMM3_Common\08_Projects_GAs\03_LAC\2024 - Snehal Yelmalle\CRIN Internal Rejection Analysis\CRIN Line rejection analysis updated.xlsx",
)


def _clean_value(value):
    if value is None:
        return ""

    if isinstance(value, datetime):
        return value.strftime("%d-%m-%Y")

    if isinstance(value, date):
        return value.strftime("%d-%m-%Y")

    return str(value).strip()


def _iso_date(value):
    if value is None or value == "":
        return ""

    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")

    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")

    text = str(value).strip()

    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass

    return ""


def _as_number(value):
    if value in (None, ""):
        return 0.0

    try:
        return float(value)

    except (TypeError, ValueError):

        try:
            return float(
                str(value).replace(",", "").strip()
            )

        except (TypeError, ValueError):
            return 0.0


def _display_number(value):
    number = _as_number(value)

    if number.is_integer():
        return int(number)

    return round(number, 2)


def load_current_excel():
    """Load the master Particle Summary sheet fresh for every request."""

    empty = {
        "headers": EXCEL_HEADERS.copy(),
        "rows": [],
        "records": [],
        "total_rows": 0,
        "error": "",
    }

    if not os.path.exists(EXCEL_FILE):
        empty["error"] = (
            f"Excel source not available: {EXCEL_FILE}"
        )
        return empty

    try:

        workbook = load_workbook(
            EXCEL_FILE,
            data_only=True,
            read_only=True
        )

        if EXCEL_SHEET not in workbook.sheetnames:

            workbook.close()

            empty["error"] = (
                f"Sheet '{EXCEL_SHEET}' was not found in the workbook."
            )

            return empty

        sheet = workbook[EXCEL_SHEET]

        raw_rows = sheet.iter_rows(values_only=True)

        source_header = next(raw_rows, None)

        if source_header is None:

            workbook.close()

            return empty

        header_map = {}

        for idx, value in enumerate(source_header):
            header_map[_clean_value(value)] = idx

        missing = [
            h for h in EXCEL_HEADERS
            if h not in header_map
        ]

        if missing:

            workbook.close()

            empty["error"] = (
                "Missing required Excel columns: "
                + ", ".join(missing)
            )

            return empty

        records = []
        rows = []

        for raw in raw_rows:

            values = []

            for header in EXCEL_HEADERS:

                idx = header_map[header]

                values.append(
                    raw[idx] if idx < len(raw) else None
                )

            if not any(_clean_value(v) for v in values):
                continue

            record = {
                header: _clean_value(values[i])
                for i, header in enumerate(EXCEL_HEADERS)
            }

            record["_date_iso"] = _iso_date(
                values[0]
            )

            record["_rejection_date_iso"] = _iso_date(
                values[11]
            )

            record["_parts"] = _as_number(
                values[6]
            )

            record["_total_particle"] = _as_number(
                values[10]
            )

            # Business rule:
            # Total particle found = 0 -> OK
            # Total particle found > 0 -> NOK
            # Total particle found blank -> PENDING

            particle_value = values[10]

            if particle_value in (None, ""):
                record["status"] = "PENDING"
            else:
                total_particle = _as_number(particle_value)

                if total_particle == 0:
                    record["status"] = "OK"
                else:
                    record["status"] = "NOK"

            record["status_label"] = record["status"]
            row = record.copy()

            row.pop("_date_iso", None)
            row.pop("_rejection_date_iso", None)
            row.pop("_parts", None)
            row.pop("_total_particle", None)
            row.pop("status_label", None)

            rows.append(
                [record[h] for h in EXCEL_HEADERS]
            )

            records.append(record)

        workbook.close()

        return {
            "headers": EXCEL_HEADERS.copy(),
            "rows": rows,
            "records": records,
            "total_rows": len(records),
            "error": "",
        }

    except Exception as exc:

        try:
            workbook.close()
        except Exception:
            pass

        empty["error"] = (
            f"Unable to read Excel source: {exc}"
        )

        return empty


def filter_excel_records(
    records,
    from_date=None,
    to_date=None,
    shift=None,
    auditor=None,
    analysis_field=None,
    analysis_value=None,
    status=None,
    rejection_shift=None,
    station=None,
    customer=None,
    record_type=None,
    metallic=None,
    da_fresh=None,
    chemistry=None,
    size=None,
    search=None
):

    result = []

    for record in records:

        if (
            from_date
            and record.get("_date_iso", "") < from_date
        ):
            continue

        if (
            to_date
            and record.get("_date_iso", "") > to_date
        ):
            continue

        if (
            shift
            and record.get("Shift") != shift
        ):
            continue

        if (
            auditor
            and record.get("Auditor") != auditor
        ):
            continue

        if (
            analysis_field
            and analysis_value
            and record.get(analysis_field) != analysis_value
        ):
            continue

        if (
            status
            and status != "All"
            and record.get("status") != status
        ):
            continue

        if (
            rejection_shift
            and record.get("Rejection Shift") != rejection_shift
        ):
            continue

        if (
            station
            and record.get("Station") != station
        ):
            continue

        if (
            customer
            and record.get("Customer") != customer
        ):
            continue

        if (
            record_type
            and record.get("Type") != record_type
        ):
            continue

        if (
            metallic
            and record.get("Metallic/Non metallic") != metallic
        ):
            continue

        if (
            da_fresh
            and record.get("DA/Fresh") != da_fresh
        ):
            continue

        if (
            chemistry
            and record.get("Chemistry") != chemistry
        ):
            continue

        if (
            size
            and record.get("Size") != size
        ):
            continue

        if search:

            needle = search.strip().lower()

            haystack = " ".join(
                str(record.get(h, ""))
                for h in EXCEL_HEADERS
            ).lower()

            if needle not in haystack:
                continue

        result.append(record)

    return result

SHIFT_ORDER = ["1st", "Gen", "2nd", "3rd"]
def excel_filter_options(records):
    fields = [
        "Shift",
        "Auditor",
        "Type",
        "Customer",
        "Station",
        "Rejection Shift",
        "Metallic/Non metallic",
        "DA/Fresh",
        "Chemistry",
        "Size",
    ]

    options = {
        field: sorted({
            str(r.get(field, "")).strip()
            for r in records
            if str(r.get(field, "")).strip()
        })
        for field in fields
    }

    # Required shift order
    options["Shift"] = (
        [shift for shift in SHIFT_ORDER if shift in options["Shift"]]
        + [
            shift
            for shift in options["Shift"]
            if shift not in SHIFT_ORDER
        ]
    )

    return options


def excel_metrics(records):
    total_records = len(records)

    ok_records = sum(
        1
        for r in records
        if r.get("status") == "OK"
    )

    pending_records = sum(
        1
        for r in records
        if r.get("status") == "PENDING"
    )

    nok_records = sum(
        1
        for r in records
        if r.get("status") == "NOK"
    )

    parts_checked = sum(
        r.get("_parts", 0)
        for r in records
    )

    total_particles = sum(
        r.get("_total_particle", 0)
        for r in records
    )

    nok_rate = round(
        (nok_records / parts_checked) * 100,
        2
    ) if parts_checked else 0

    return {
        "total_records": total_records,
        "ok_records": ok_records,
        "pending_records": pending_records,
        "nok_records": nok_records,
        "parts_checked": _display_number(parts_checked),
        "total_particles": _display_number(total_particles),
        "nok_rate": nok_rate,
    }

def _group_counts(records, field):

    grouped = {}

    for record in records:

        key = record.get(field) or "Unknown"

        item = grouped.setdefault(
            key,
            {
                "parts": 0,
                "nok": 0
            }
        )

        item["parts"] += record.get(
            "_parts",
            0
        )

        if record.get("status") == "NOK":
            item["nok"] += 1

    keys = list(grouped.keys())

    return (
        keys,

        [
            _display_number(
                grouped[k]["parts"]
            )
            for k in keys
        ],

        [
            grouped[k]["nok"]
            for k in keys
        ],

        [
            round(
                (
                    grouped[k]["nok"]
                    /
                    grouped[k]["parts"]
                    *
                    100
                ),
                2
            )
            if grouped[k]["parts"]
            else 0
            for k in keys
        ],

        grouped,
    )


# ================= DASHBOARD =================

@app.get(
    "/dashboard",
    response_class=HTMLResponse
)
def dashboard(
    request: Request,
    from_date: str = None,
    to_date: str = None,
    shift: str = None,
    auditor: str = None,
    type: str = None
):

    try:

        user = get_current_user(request)

    except Exception:

        return RedirectResponse("/login")

    data = load_current_excel()

    records = filter_excel_records(
        data["records"],
        from_date,
        to_date,
        shift=shift,
        auditor=auditor,
        record_type=type
    )

    metrics = excel_metrics(records)

    options = excel_filter_options(
        data["records"]
    )


    # ================= MONTHLY RECORDS VS NOK =================

    monthly = {}

    for record in records:

        date_value = (
            record.get("_date_iso")
            or "Unknown"
        )

        if date_value != "Unknown":

            month_key = str(
                date_value
            )[:7]

        else:

            month_key = "Unknown"

        monthly.setdefault(
    month_key,
    {
        "records": 0,
        "ok": 0,
        "nok": 0,
        "pending": 0
    }
)

       # Count total records by status

        monthly[month_key]["records"] += 1

        if record.get("status") == "NOK":
            monthly[month_key]["nok"] += 1

        elif record.get("status") == "PENDING":
            monthly[month_key]["pending"] += 1

        elif record.get("status") == "OK":
            monthly[month_key]["ok"] += 1


    # ================= SORT MONTHS =================

    monthly_keys = sorted(
        monthly.keys()
    )


    # ================= MONTH LABELS =================

    month_names = [
        "Jan",
        "Feb",
        "Mar",
        "Apr",
        "May",
        "Jun",
        "Jul",
        "Aug",
        "Sep",
        "Oct",
        "Nov",
        "Dec"
    ]

    monthly_labels = []

    for key in monthly_keys:

        if key == "Unknown":

            monthly_labels.append(
                "Unknown"
            )

        else:

            year = key[:4]

            month = key[5:7]

            try:

                month_name = month_names[
                    int(month) - 1
                ]

                monthly_labels.append(
                    f"{month_name} {year}"
                )

            except (
                ValueError,
                IndexError
            ):

                monthly_labels.append(
                    key
                )


    # ================= MONTHLY RECORD VALUES =================

    monthly_records = [
        monthly[key]["records"]
        for key in monthly_keys
    ]


    # ================= MONTHLY NOK VALUES =================

    monthly_nok = [
        monthly[key]["nok"]
        for key in monthly_keys
    ]
    monthly_ok = [
    monthly[key]["ok"]
    for key in monthly_keys
]
    


    # ================= SHIFT-WISE NOK =================

    shift_nok = {}

    for record in records:

        key = (
            record.get("Shift")
            or "Unknown"
        )

        shift_nok[key] = (
            shift_nok.get(key, 0)
            +
            (
                1
                if record.get("status") == "NOK"
                else 0
            )
        )


    # ================= STATION-WISE NOK =================

        # ================= STATION-WISE NOK =================

    station_nok = {}

    for record in records:
        if record.get("status") != "NOK":
            continue

        key = (
            record.get("Station")
            or "Unknown"
        )

        station_nok[key] = station_nok.get(key, 0) + 1


    # ================= AUDITOR COUNTS =================

    auditor_counts = {}

    for record in records:

        key = (
            record.get("Auditor")
            or "Unknown"
        )

        auditor_counts[key] = (
            auditor_counts.get(key, 0)
            + 1
        )


    auditor_counts = dict(
        sorted(
            auditor_counts.items(),
            key=lambda item: (
                -item[1],
                item[0]
            )
        )
    )


    # ================= RECENT NOK =================

    recent_nok = [
        record
        for record in records
        if record.get("status") == "NOK"
    ]

    recent_nok.sort(
        key=lambda record: (
            record.get(
                "_rejection_date_iso"
            )
            or record.get(
                "_date_iso"
            )
            or ""
        ),
        reverse=True
    )


    # ================= SEND DATA TO DASHBOARD =================

    return templates.TemplateResponse(
        "dashboard.html",
        {
            "request": request,

            "user": user,

            "current_date": date.today().strftime(
                "%d %b %Y"
            ),

            **metrics,

            "filters": {
                "from_date": from_date,
                "to_date": to_date,
                "shift": shift,
                "auditor": auditor,
                "type": type
            },

            "shifts": options["Shift"],

            "auditors": options["Auditor"],

            "types": options["Type"],

            # Kept with the existing names so
            # dashboard.js does not need another
            # backend variable-name change.

            "daily_labels": monthly_labels,

            "daily_records": monthly_records,

            "daily_nok": monthly_nok,
            "daily_ok": monthly_ok,

            "shift_labels": list(
                shift_nok.keys()
            ),

            "shift_nok": list(
                shift_nok.values()
            ),

            "station_labels": list(
                station_nok.keys()
            ),

            "station_nok": list(
                station_nok.values()
            ),

            "auditor_counts": auditor_counts,

            "recent_nok": recent_nok[:8],

            "excel_headers": data["headers"],

            "excel_rows": [
                [
                    record[h]
                    for h in EXCEL_HEADERS
                ]
                for record in data["records"]
            ],

            "excel_total_rows": (
                data["total_rows"]
            ),

            "source_error": data["error"],

            "show_header": True
        }
    )


# ================= RECORD ENTRY =================

# Record creation/editing is intentionally disabled
# in the Excel-driven WSA version.
# The live Bosch workbook remains the single source
# for record data.


# ================= VIEW RECORDS =================

@app.get(
    "/view",
    response_class=HTMLResponse
)
def view_records(
    request: Request,
    from_date: str = None,
    to_date: str = None,
    shift: str = None,
    auditor: str = None,
    type: str = None,
    status: str = None,
    sort_by: str = "date",
    order: str = "desc",
    search: str = None
):

    try:

        user = get_current_user(request)

    except Exception:

        return RedirectResponse("/login")

    data = load_current_excel()

    records = filter_excel_records(
        data["records"],
        from_date,
        to_date,
        shift=shift,
        auditor=auditor,
        record_type=type,
        status=status,
        search=search
    )

    sort_map = {
        "date": "_date_iso",
        "rejection_date": "_rejection_date_iso",
        "parts": "_parts",
        "particles": "_total_particle",
        "auditor": "Auditor",
        "shift": "Shift",
        "station": "Station",
    }

    key = sort_map.get(
        sort_by,
        "_date_iso"
    )

    records.sort(
        key=lambda r: (
            r.get(key) or ""
        ),
        reverse=(
            order == "desc"
        )
    )

    metrics = excel_metrics(
        records
    )

    options = excel_filter_options(
        data["records"]
    )

    return templates.TemplateResponse(
        "view_records.html",
        {
            "request": request,

            "user": user,

            "records": records,

            **metrics,

            "options": options,

            "filters": {
                "from_date": from_date,
                "to_date": to_date,
                "shift": shift,
                "auditor": auditor,
                "type": type,
                "status": status,
                "sort_by": sort_by,
                "order": order,
                "search": search,
            },

            "headers": EXCEL_HEADERS,

            "source_error": data["error"],

            "show_header": True
        }
    )


# ================= REPORTS & EXPORT =================

@app.get(
    "/download",
    response_class=HTMLResponse
)
def download_page(request: Request):

    try:

        get_current_user(request)

    except Exception:

        return RedirectResponse("/login")

    data = load_current_excel()

    return templates.TemplateResponse(
        "download.html",
        {
            "request": request,

            "user": get_current_user(
                request
            ),

            "total_rows": data["total_rows"],

            "source_error": data["error"],

            "show_header": True
        }
    )


@app.post("/download")
def download_csv(
    request: Request,
    from_date: str = Form(None),
    to_date: str = Form(None),
    status: str = Form("All")
):

    try:

        get_current_user(request)

    except Exception:

        return RedirectResponse("/login")

    data = load_current_excel()

    records = filter_excel_records(
        data["records"],
        from_date,
        to_date,
        status=status
    )

    output = io.StringIO()

    writer = csv.writer(
        output
    )

    writer.writerow(
        EXCEL_HEADERS + ["Status"]
    )

    for record in records:

        writer.writerow(
            [
                record.get(
                    h,
                    ""
                )
                for h in EXCEL_HEADERS
            ]
            +
            [
                record.get(
                    "status",
                    ""
                )
            ]
        )

    output.seek(0)

    filename = (
        "wsa_excel_records.csv"
    )

    return StreamingResponse(
        iter(
            [
                output.getvalue()
            ]
        ),
        media_type="text/csv; charset=utf-8",

        headers={
            "Content-Disposition":
                f"attachment; filename={filename}"
        }
    )


# ================= ANALYTICS =================

ANALYSIS_FIELDS = [
    "Shift",
    "Auditor",
    "Type",
    "Customer",
    "Station",
    "Rejection Shift",
    "Metallic/Non metallic",
    "DA/Fresh",
    "Chemistry",
    "Size"
]


@app.get(
    "/analytics",
    response_class=HTMLResponse
)
def analytics_page(
    request: Request,
    analysis_type: str = "Shift",
    from_date: str = None,
    to_date: str = None,
    shift: str = None,
    status: str = None
):

    try:

        user = get_current_user(request)

    except Exception:

        return RedirectResponse("/login")

    if analysis_type not in ANALYSIS_FIELDS:

        analysis_type = "Shift"

    data = load_current_excel()

    records = filter_excel_records(
        data["records"],
        from_date,
        to_date,
        shift=shift,
        status=status
    )

    labels, parts, nok, nok_rates, grouped = (
        _group_counts(
            records,
            analysis_type
        )
    )

    metrics = excel_metrics(
        records
    )

    options = excel_filter_options(
        data["records"]
    )

    breakdown = []

    for label in labels:

        breakdown.append(
            {
                "group": label,

                "parts": grouped[label]["parts"],

                "nok": grouped[label]["nok"],

                "nok_rate": round(
                    (
                        grouped[label]["nok"]
                        /
                        grouped[label]["parts"]
                        *
                        100
                    ),
                    2
                )
                if grouped[label]["parts"]
                else 0,
            }
        )

    breakdown.sort(
        key=lambda item: (
            -item["nok"],
            str(item["group"])
        )
    )

    return templates.TemplateResponse(
        "analytics.html",
        {
            "request": request,

            "user": user,

            "analysis_fields": ANALYSIS_FIELDS,

            "analysis_type": analysis_type,

            "labels": labels,

            "parts": parts,

            "nok": nok,

            "nok_rates": nok_rates,

            "breakdown": breakdown,

            **metrics,

            "options": options,

            "filters": {
                "from_date": from_date,
                "to_date": to_date,
                "shift": shift,
                "status": status
            },

            "source_error": data["error"],

            "show_header": True
        }
    )


@app.post(
    "/analytics",
    response_class=HTMLResponse
)
def analytics_data(
    request: Request,
    analysis_type: str = Form(...),
    from_date: str = Form(None),
    to_date: str = Form(None),
    shift: str = Form(None),
    status: str = Form(None)
):

    params = [
        f"analysis_type={analysis_type}"
    ]

    if from_date:
        params.append(
            f"from_date={from_date}"
        )

    if to_date:
        params.append(
            f"to_date={to_date}"
        )

    if shift:
        params.append(
            f"shift={shift}"
        )

    if status:
        params.append(
            f"status={status}"
        )

    return RedirectResponse(
        "/analytics?" + "&".join(params),
        status_code=303
    )


# ================= OBSERVATIONS =================

@app.get(
    "/observations",
    response_class=HTMLResponse
)
def observations(
    request: Request,
    from_date: str = None,
    to_date: str = None,
    shift: str = None,
    auditor: str = None,
    type: str = None,
    status: str = "NOK",
    rejection_shift: str = None,
    station: str = None,
    customer: str = None,
    metallic: str = None,
    da_fresh: str = None,
    chemistry: str = None,
    size: str = None,
    search: str = None
):

    try:

        user = get_current_user(request)

    except Exception:

        return RedirectResponse("/login")

    data = load_current_excel()

    records = filter_excel_records(
        data["records"],
        from_date,
        to_date,
        shift=shift,
        auditor=auditor,
        record_type=type,
        status=status,
        rejection_shift=rejection_shift,
        station=station,
        customer=customer,
        metallic=metallic,
        da_fresh=da_fresh,
        chemistry=chemistry,
        size=size,
        search=search
    )

    records.sort(
        key=lambda r: (
            r.get(
                "_rejection_date_iso"
            )
            or r.get(
                "_date_iso"
            )
            or ""
        ),
        reverse=True
    )

    metrics = excel_metrics(
        records
    )

    options = excel_filter_options(
        data["records"]
    )

    particle_rows = sum(
        1
        for r in records
        if r.get(
            "_total_particle",
            0
        ) > 0
    )

    return templates.TemplateResponse(
        "observations.html",
        {
            "request": request,

            "user": user,

            "records": records,

            "headers": EXCEL_HEADERS,

            "particle_rows": particle_rows,

            **metrics,

            "options": options,

            "filters": {
                "from_date": from_date,
                "to_date": to_date,
                "shift": shift,
                "auditor": auditor,
                "type": type,
                "status": status,
                "rejection_shift": rejection_shift,
                "station": station,
                "customer": customer,
                "metallic": metallic,
                "da_fresh": da_fresh,
                "chemistry": chemistry,
                "size": size,
                "search": search,
            },

            "source_error": data["error"],

            "show_header": True
        }
    )


# Legacy database CRUD/case/debug endpoints
# were removed from the Excel-driven web workflow.

# ================= ALERTS API =================

from backend.alerts import service as alert_service
from backend.alerts import schemas as alert_schemas

@app.post("/api/alerts/evaluate")
def trigger_alert_evaluation(db: Session = Depends(get_db)):
    # This endpoint manually triggers an evaluation against current Excel records
    # Usually this could be scheduled by a background worker, but for now it's an API trigger
    data = load_current_excel()
    if data["error"]:
        raise HTTPException(status_code=500, detail=data["error"])
        
    events = alert_service.evaluate_and_trigger(db, data["records"])
    return {"message": f"Evaluated {len(data['records'])} records. Triggered {len(events)} new events."}

@app.post("/api/alerts/poll", response_model=list[alert_schemas.PowerAutomatePollResponse])
def poll_pending_alerts(db: Session = Depends(get_db)):
    # Power Automate will call this to fetch pending alerts
    return alert_service.get_pending_alerts_for_power_automate(db)

@app.post("/api/alerts/mark-sent/{alert_id}")
def mark_alert_sent(alert_id: int, db: Session = Depends(get_db)):
    # Power Automate will call this after successfully sending the email
    alert_service.mark_alert_as_sent(db, alert_id)
    return {"status": "success"}

# ================= ALERTS ADMIN UI =================

@app.get("/admin/alerts", response_class=HTMLResponse)
def alert_dashboard(request: Request, db: Session = Depends(get_db)):
    try:
        user = get_current_user(request)
        if user.role.lower() != "admin":
            return RedirectResponse("/")
    except Exception:
        return RedirectResponse("/login")
        
    rules = db.query(alert_models.AlertRule).all()
    events = db.query(alert_models.AlertEvent).order_by(alert_models.AlertEvent.triggered_at.desc()).limit(50).all()
    
    return templates.TemplateResponse("alerts_dashboard.html", {
        "request": request,
        "user": user,
        "rules": rules,
        "events": events,
        "show_header": True
    })

@app.get("/admin/alerts/builder", response_class=HTMLResponse)
def alert_builder(request: Request):
    try:
        user = get_current_user(request)
        if user.role.lower() != "admin":
            return RedirectResponse("/")
    except Exception:
        return RedirectResponse("/login")
        
    return templates.TemplateResponse("alerts_builder.html", {
        "request": request,
        "user": user,
        "show_header": True
    })
