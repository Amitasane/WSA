"""
Automated Test Suite for Investigation Dashboard
Verifies data loading, KPI accuracy, Pareto distributions, filters, routes, and export endpoints.
"""

import os
import io
import csv
import openpyxl
from fastapi.testclient import TestClient

from backend.main import app
from backend.investigation_loader import (
    load_investigation_workbook,
    get_cached_investigation_data,
    reload_investigation_data,
    get_excel_file_path,
)
from backend.investigation_analytics import (
    calculate_kpis,
    filter_investigation_records,
    get_monthly_trend,
    get_pareto_defects,
    get_customer_breakdown,
    get_product_share,
    get_analytics_breakdown,
)


def test_data_loader():
    data = load_investigation_workbook()
    assert not data["error"], f"Data loader error: {data['error']}"
    assert data["total_records"] == 275, f"Expected 275 records, got {data['total_records']}"
    assert len(data["cri_records"]) == 183, f"Expected 183 CRI records, got {len(data['cri_records'])}"
    assert len(data["crin_records"]) == 92, f"Expected 92 CRIN records, got {len(data['crin_records'])}"
    assert len(data["months"]) >= 9, f"Expected >= 9 months, got {len(data['months'])}"
    assert "M&M" in data["customers"], "Customer M&M missing"
    assert "TCL" in data["customers"], "Customer TCL missing"
    print("PASS: Data loader tests")


def test_analytics_and_kpis():
    data = get_cached_investigation_data()
    records = data["records"]
    kpis = calculate_kpis(records)

    assert kpis["total_investigations"] == 275
    assert kpis["total_parts"] == 475
    assert kpis["completed"] == 198
    assert kpis["pending"] == 77
    assert kpis["defect_cases"] == 155
    assert kpis["conforming_cases"] == 43
    assert kpis["defect_rate"] > 75.0, f"Unexpected defect rate: {kpis['defect_rate']}"

    # Monthly Trend
    trend = get_monthly_trend(records)
    assert len(trend["labels"]) == 9
    assert sum(trend["total"]) == 270  # 5 records with blank/unspecified month


    # Pareto Defects
    pareto = get_pareto_defects(records)
    assert pareto["labels"][0] == "Particle at Z-Hole", f"Top defect should be Particle at Z-Hole, got {pareto['labels'][0]}"
    assert pareto["counts"][0] == 59
    assert pareto["cumulative_pct"][-1] == 100.0

    # Dimension Grouping
    for dim in ["Customer", "Defect Category", "Complaint Type", "Product Line", "Month"]:
        breakdown = get_analytics_breakdown(records, dim)
        assert len(breakdown["rows"]) > 0, f"Dimension {dim} produced 0 rows"
    print("PASS: Analytics and KPI calculation tests")


def test_fastapi_endpoints():
    client = TestClient(app)
    cookies = {"user_id": "1"}

    # 1. Login Page
    resp = client.get("/login")
    assert resp.status_code == 200
    assert "Investigation" in resp.text

    # 2. Unauthenticated Redirect
    resp = client.get("/dashboard", follow_redirects=False)
    assert resp.status_code in (302, 307)

    # 3. Authenticated Home Page
    resp = client.get("/", cookies=cookies)
    assert resp.status_code == 200
    assert "Investigation Dashboard" in resp.text

    # 4. Authenticated Dashboard Page
    resp = client.get("/dashboard", cookies=cookies)
    assert resp.status_code == 200
    assert "trendChart" in resp.text
    assert "paretoChart" in resp.text
    assert "customerChart" in resp.text
    assert "productShareChart" in resp.text

    # 5. Dashboard Filters
    resp = client.get("/dashboard?product_class=CRI&month=2026-02", cookies=cookies)
    assert resp.status_code == 200

    # 6. View Records Page
    resp = client.get("/view", cookies=cookies)
    assert resp.status_code == 200
    assert "Investigation Master Table" in resp.text
    assert "CRI.I.26.01" in resp.text

    # 7. View Records Search & Pagination
    resp = client.get("/view?search=0445120568&page=1&page_size=25", cookies=cookies)
    assert resp.status_code == 200
    assert "0445120568" in resp.text

    # 8. Analytics Page
    resp = client.get("/analytics", cookies=cookies)
    assert resp.status_code == 200
    assert "Statistical Summary" in resp.text

    # 9. Observations Page
    resp = client.get("/observations", cookies=cookies)
    assert resp.status_code == 200
    assert "Quality Observations" in resp.text

    # 10. Observations Search
    resp = client.get("/observations?search=Z+hole", cookies=cookies)
    assert resp.status_code == 200
    assert "Z hole" in resp.text or "Z-hole" in resp.text

    # 11. Reports & Export Page
    resp = client.get("/download", cookies=cookies)
    assert resp.status_code == 200

    # 12. CSV Export
    resp = client.post("/download", data={"format": "csv", "product_class": "All"}, cookies=cookies)
    assert resp.status_code == 200
    assert "text/csv" in resp.headers["content-type"]
    reader = list(csv.reader(io.StringIO(resp.text)))
    assert len(reader) == 276, f"Expected 276 CSV rows, got {len(reader)}"

    # 13. Excel Export
    resp = client.post("/download", data={"format": "xlsx", "product_class": "All"}, cookies=cookies)
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws.max_row == 276, f"Expected 276 Excel rows, got {ws.max_row}"

    # 14. Refresh Data Endpoint
    resp = client.get("/refresh-data", cookies=cookies, follow_redirects=False)
    assert resp.status_code == 302

    print("PASS: All FastAPI endpoint tests")


if __name__ == "__main__":
    test_data_loader()
    test_analytics_and_kpis()
    test_fastapi_endpoints()
    print("\n=============================================")
    print("ALL TESTS PASSED SUCCESSFULLY! (100% SUCCESS)")
    print("=============================================")
